"""Writes validated records to an .xlsx workbook for manual human review of
generated Q&A pairs before they're trusted for QLoRA training."""
from __future__ import annotations

import logging
from pathlib import Path

from openpyxl import Workbook
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from quality_gate import ValidatedRecord

logger = logging.getLogger(__name__)

_HEADERS = [
    "Review", "Notes", "Type", "Question", "Answer",
    "System Prompt", "Source Doc", "Page Start", "Page End", "Chunk ID",
]
_COLUMN_WIDTHS = [10, 28, 8, 45, 45, 35, 22, 10, 10, 20]
_REVIEW_OPTIONS = ["Keep", "Reject", "Needs Fix"]

_HEADER_FILL = PatternFill("solid", fgColor="2F2F2F")
_KEEP_FILL = PatternFill("solid", fgColor="1F4D2B")
_REJECT_FILL = PatternFill("solid", fgColor="5C1F1F")
_NEEDS_FIX_FILL = PatternFill("solid", fgColor="5C4A1F")


def _write_sheet(ws, records: list[ValidatedRecord]) -> None:
    ws.append(_HEADERS)
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = _HEADER_FILL

    for r in records:
        ws.append([
            "", "", "Table" if r.is_table else "Text",
            r.question, r.answer, r.system,
            r.source_doc, r.page_start, r.page_end, r.chunk_id,
        ])

    last_row = ws.max_row
    last_col_letter = get_column_letter(len(_HEADERS))

    if last_row > 1:
        dv = DataValidation(
            type="list",
            formula1='"{}"'.format(",".join(_REVIEW_OPTIONS)),
            allow_blank=True,
        )
        ws.add_data_validation(dv)
        dv.add(f"A2:A{last_row}")

        row_range = f"A2:{last_col_letter}{last_row}"
        ws.conditional_formatting.add(row_range, FormulaRule(formula=['$A2="Keep"'], fill=_KEEP_FILL))
        ws.conditional_formatting.add(row_range, FormulaRule(formula=['$A2="Reject"'], fill=_REJECT_FILL))
        ws.conditional_formatting.add(row_range, FormulaRule(formula=['$A2="Needs Fix"'], fill=_NEEDS_FIX_FILL))

    for idx, width in enumerate(_COLUMN_WIDTHS, start=1):
        ws.column_dimensions[get_column_letter(idx)].width = width
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = cell.alignment.copy(wrap_text=True, vertical="top")

    ws.freeze_panes = "A2"


def write_review_workbook(splits: dict[str, list[ValidatedRecord]], output_dir: str) -> str:
    wb = Workbook()
    wb.remove(wb.active)

    for split_name, records in splits.items():
        ws = wb.create_sheet(title=split_name)
        _write_sheet(ws, records)

    out_dir = Path(output_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    out_path = out_dir / "review.xlsx"
    wb.save(out_path)
    logger.info("Wrote human-review workbook to %s", out_path)
    return str(out_path)
